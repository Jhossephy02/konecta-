from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.pagesizes import A4
root=Path(__file__).resolve().parent.parent
out=root/'public/templates';out.mkdir(parents=True,exist_ok=True)
intro='Reemplaza los ejemplos con tus preguntas y respuestas. Conserva las etiquetas PREGUNTA y RESPUESTA y las letras A), B), C). Puedes usar de 2 a 6 opciones por pregunta y de 1 a 10 preguntas. Cada pregunta y cada opción debe ocupar un solo párrafo. Guarda como DOCX o exporta a PDF con texto seleccionable; no uses escaneos. Revisa la importación antes de publicar el curso.'
blocks=[['PREGUNTA: ¿Dónde debes revisar tus capacitaciones asignadas?','A) En Mi aprendizaje','B) En Cumpleaños','C) En una cuenta de otra persona','RESPUESTA: A'],['PREGUNTA: ¿Qué debes hacer al terminar de usar una tablet compartida?','A) Dejar abierta la sesión','B) Cerrar sesión y dejar la tablet','C) Compartir tu contraseña','RESPUESTA: B']]
d=Document();section=d.sections[0];section.top_margin=section.bottom_margin=Inches(.8)
normal=d.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(11)
for name in ['Title','Heading 1','Heading 2']:d.styles[name].font.color.rgb=RGBColor(0,0,0)
d.add_paragraph('Plantilla de evaluación de cursos','Title');d.add_paragraph(intro);d.add_paragraph('Preguntas de ejemplo','Heading 1')
for block in blocks:
 for line in block:d.add_paragraph(line)
 d.add_paragraph('')
for element in d.styles.element.xpath('.//w:pBdr'):
 element.getparent().remove(element)
d.save(out/'evaluacion.docx')
styles=getSampleStyleSheet();styles['Normal'].fontSize=11;styles['Normal'].leading=16
story=[Paragraph('Plantilla de evaluación de cursos',styles['Title']),Spacer(1,14),Paragraph(intro,styles['Normal']),Spacer(1,22),Paragraph('Preguntas de ejemplo',styles['Heading1'])]
for block in blocks:
 for line in block:story.append(Paragraph(line,styles['Normal']))
 story.append(Spacer(1,20))
SimpleDocTemplate(str(out/'evaluacion.pdf'),pagesize=A4,rightMargin=52,leftMargin=52,topMargin=48,bottomMargin=48).build(story)
print('Plantillas generadas')
